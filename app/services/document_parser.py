from __future__ import annotations

import re
import shutil
import os
import subprocess
from pathlib import Path
from typing import Iterable

from ..schemas import DocumentSegment, ParsedDocument, SourceLocation

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _segment(
    text: str,
    source_id: str,
    label: str,
    *,
    page: int | None = None,
    paragraph: int | None = None,
    table: int | None = None,
    row: int | None = None,
    column: int | None = None,
    line_start: int | None = None,
    line_end: int | None = None,
    start_offset: int,
) -> DocumentSegment:
    cleaned = text.strip()
    return DocumentSegment(
        source_id=source_id,
        text=cleaned,
        location=SourceLocation(
            source_id=source_id,
            page=page,
            paragraph=paragraph,
            table=table,
            row=row,
            column=column,
            line_start=line_start,
            line_end=line_end,
            label=label,
        ),
        start_offset=start_offset,
        end_offset=start_offset + len(cleaned),
    )


def _make_line_segments(text: str, filename: str) -> list[DocumentSegment]:
    segments: list[DocumentSegment] = []
    offset = 0
    for line_number, line in enumerate(text.splitlines(), start=1):
        if line.strip():
            source_id = f"line-{line_number}"
            segments.append(
                _segment(
                    line,
                    source_id,
                    f"第 {line_number} 行",
                    line_start=line_number,
                    line_end=line_number,
                    start_offset=offset,
                )
            )
        offset += len(line) + 1
    return segments


def _parse_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    segments = _make_line_segments(text, path.name)
    warnings = [] if segments else ["文件没有可提取的文字内容。"]
    return ParsedDocument(
        filename=path.name,
        media_type="text/plain",
        segments=segments,
        warnings=warnings,
        total_characters=len(text),
    )


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请先运行 run.bat 安装依赖。") from exc

    document = Document(path)
    segments: list[DocumentSegment] = []
    warnings: list[str] = []
    offset = 0
    paragraph_number = 0
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            source_id = f"paragraph-{paragraph_number}"
            segments.append(
                _segment(
                    paragraph.text,
                    source_id,
                    f"第 {paragraph_number} 个段落",
                    paragraph=paragraph_number,
                    start_offset=offset,
                )
            )
            offset += len(paragraph.text) + 1

    for section_number, section in enumerate(document.sections, start=1):
        for area_name, container in (("页眉", section.header), ("页脚", section.footer)):
            for paragraph_number, paragraph in enumerate(container.paragraphs, start=1):
                if not paragraph.text.strip():
                    continue
                source_id = f"{area_name}-{section_number}-paragraph-{paragraph_number}"
                if any(existing.source_id == source_id for existing in segments):
                    continue
                segments.append(
                    _segment(
                        paragraph.text,
                        source_id,
                        f"第 {section_number} 节{area_name}，第 {paragraph_number} 个段落",
                        paragraph=paragraph_number,
                        start_offset=offset,
                    )
                )
                offset += len(paragraph.text) + 1

    for table_number, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        headers = rows[0]
        seen_rows: set[tuple[str, ...]] = set()
        for row_number, row_cells in enumerate(rows, start=1):
            normalized_cells: list[str] = []
            for cell_text in row_cells:
                normalized_cells.append(cell_text)
            if not any(normalized_cells):
                continue
            row_key = tuple(normalized_cells)
            if row_key in seen_rows:
                continue
            seen_rows.add(row_key)
            if row_number == 1:
                text = "表头：" + " | ".join(normalized_cells)
            else:
                pairs = []
                for column_number, cell_text in enumerate(normalized_cells, start=1):
                    header = headers[column_number - 1] if column_number <= len(headers) and headers[column_number - 1] else f"第{column_number}列"
                    pairs.append(f"{header}：{cell_text}")
                text = f"表头：{' | '.join(headers)}\n第 {row_number} 行：" + " | ".join(pairs)
            source_id = f"table-{table_number}-r{row_number}"
            segments.append(
                _segment(
                    text,
                    source_id,
                    f"第 {table_number} 个表格，第 {row_number} 行（含表头上下文）",
                    table=table_number,
                    row=row_number,
                    start_offset=offset,
                )
            )
            offset += len(text) + 1

    if not segments:
        warnings.append("DOCX 没有可提取的文字内容。")
    return ParsedDocument(
        filename=path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        segments=segments,
        warnings=warnings,
        total_characters=offset,
    )


def _split_pdf_pages(text: str) -> Iterable[tuple[int, str]]:
    pages = text.split("\f")
    for page_number, page_text in enumerate(pages, start=1):
        if page_text.strip():
            yield page_number, page_text


def _parse_pdf(path: Path) -> ParsedDocument:
    text = ""
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz

        with fitz.open(path) as pdf:
            pages = [page.get_text("text") for page in pdf]
        text = "\f".join(pages)
    except ImportError:
        executable_value = os.getenv("CHECKER_PDFTOTEXT_PATH")
        executable = executable_value or shutil.which("pdftotext")
        if not executable:
            raise RuntimeError("缺少 PyMuPDF，且找不到系统 pdftotext，无法解析 PDF。")
        completed = subprocess.run(
            [executable, "-enc", "UTF-8", "-layout", str(path), "-"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=90,
            check=False,
        )
        if completed.returncode != 0:
            raise RuntimeError("PDF 文本提取失败，请确认文件未损坏。")
        text = completed.stdout

    segments: list[DocumentSegment] = []
    offset = 0
    for page_number, page_text in _split_pdf_pages(text):
        lines = page_text.splitlines()
        nonempty = [(line_number, line) for line_number, line in enumerate(lines, start=1) if line.strip()]
        for block_start in range(0, len(nonempty), 5):
            block = nonempty[block_start:block_start + 5]
            first_line, _ = block[0]
            last_line, _ = block[-1]
            block_text = "\n".join(line for _, line in block)
            source_id = f"page-{page_number}-lines-{first_line}-{last_line}"
            segments.append(
                _segment(
                    block_text,
                    source_id,
                    f"第 {page_number} 页，第 {first_line}-{last_line} 行（保留相邻上下文）",
                    page=page_number,
                    line_start=first_line,
                    line_end=last_line,
                    start_offset=offset + sum(len(line) + 1 for line in lines[:first_line - 1]),
                )
            )
        offset += sum(len(line) + 1 for line in lines) + 1

    warnings = []
    if not segments:
        warnings.append("PDF 没有可提取的文字层，可能是扫描件，请先进行 OCR。")
    elif len("".join(segment.text for segment in segments)) < 50:
        warnings.append("PDF 可提取文字很少，可能是扫描件或文字层不完整，请人工核验。")
    return ParsedDocument(
        filename=path.name,
        media_type="application/pdf",
        segments=segments,
        warnings=warnings,
        total_characters=offset,
    )


def parse_document(path: str | Path) -> ParsedDocument:
    source = Path(path)
    extension = source.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("暂不支持此文件格式，请上传 PDF、DOCX、TXT 或 MD。")
    if extension == ".pdf":
        return _parse_pdf(source)
    if extension == ".docx":
        return _parse_docx(source)
    return _parse_text(source)


def normalize_whitespace(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text).strip()
