import tempfile
import unittest
from pathlib import Path

from docx import Document

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from app.services.document_parser import parse_document
from app.services.chunker import build_chunks


class ParserTests(unittest.TestCase):
    def test_text_keeps_line_locations(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notice.txt"
            path.write_text("资格条件：具有营业执照。\n\n投标截止时间：2026年10月20日。\n", encoding="utf-8")
            document = parse_document(path)
        self.assertEqual(len(document.segments), 2)
        self.assertEqual(document.segments[0].source_id, "line-1")
        self.assertEqual(document.segments[1].location.line_start, 3)
        self.assertIn("截止时间", document.segments[1].text)

    def test_chunk_preserves_source_segments(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notice.md"
            path.write_text("资格条件\n" + "材料要求\n" * 100, encoding="utf-8")
            document = parse_document(path)
        chunks = build_chunks(document.segments, max_chars=500)
        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0].segments[0].source_id, "line-1")

    def test_docx_table_keeps_header_context(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "截止时间"
            table.cell(1, 0).text = "投标文件"
            table.cell(1, 1).text = "2026年10月20日"
            document.save(path)
            parsed = parse_document(path)
        row = next(segment for segment in parsed.segments if segment.source_id == "table-1-r2")
        self.assertIn("表头：项目 | 截止时间", row.text)
        self.assertIn("截止时间：2026年10月20日", row.text)

    def test_pdf_keeps_adjacent_line_context(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notice.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Item    Deadline\nBid file  2026-10-20\nNote    Beijing time")
            document.save(path)
            document.close()
            parsed = parse_document(path)
        self.assertEqual(len(parsed.segments), 1)
        segment = parsed.segments[0]
        self.assertIn("Item", segment.text)
        self.assertIn("2026-10-20", segment.text)
        self.assertEqual(segment.location.page, 1)
        self.assertEqual(segment.location.line_start, 1)
        self.assertEqual(segment.location.line_end, 3)

    def test_long_line_is_split_without_exceeding_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "long.txt"
            path.write_text("资格条件：" + "字" * 1400, encoding="utf-8")
            document = parse_document(path)
        chunks = build_chunks(document.segments, max_chars=500, max_chunks=10)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.text) <= 500 for chunk in chunks))
        source_ids = [segment.source_id for chunk in chunks for segment in chunk.segments]
        self.assertEqual(len(source_ids), len(set(source_ids)))
        self.assertTrue(all(source_id.startswith("line-1-part-") for source_id in source_ids))
        offsets = [segment.start_offset for chunk in chunks for segment in chunk.segments]
        self.assertEqual(offsets, sorted(offsets))


if __name__ == "__main__":
    unittest.main()
